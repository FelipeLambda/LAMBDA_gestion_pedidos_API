import csv
from io import BytesIO
from datetime import datetime
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ReporteExcelGenerator:

    def __init__(self, titulo, headers, datos):
        self.titulo = titulo
        self.headers = headers
        self.datos = datos
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = titulo[:31]

    def _aplicar_estilos_header(self):
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)

        for col_num, header in enumerate(self.headers, 1):
            cell = self.worksheet.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

    def _escribir_datos(self):
        for row_num, fila in enumerate(self.datos, 2):
            for col_num, valor in enumerate(fila, 1):
                cell = self.worksheet.cell(row=row_num, column=col_num)
                cell.value = valor
                cell.alignment = Alignment(vertical='center')

    def _ajustar_anchos(self):
        for col in range(1, len(self.headers) + 1):
            max_length = len(str(self.headers[col - 1]))
            for row in range(2, len(self.datos) + 2):
                cell_value = self.worksheet.cell(row=row, column=col).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            adjusted_width = min(max_length + 2, 50)
            self.worksheet.column_dimensions[get_column_letter(col)].width = adjusted_width

    def generar(self):
        self._aplicar_estilos_header()
        self._escribir_datos()
        self._ajustar_anchos()

        output = BytesIO()
        self.workbook.save(output)
        output.seek(0)

        filename = f"{self.titulo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class ReporteCSVGenerator:

    def __init__(self, titulo, headers, datos):
        self.titulo = titulo
        self.headers = headers
        self.datos = datos

    def generar(self):
        filename = f"{self.titulo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        response.write('\ufeff')

        writer = csv.writer(response)
        writer.writerow(self.headers)

        for fila in self.datos:
            writer.writerow(fila)

        return response
